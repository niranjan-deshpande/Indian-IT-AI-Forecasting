"""Tech Mahindra (techm) extraction: fact sheets (PDF text), data sheets (xlsx), consolidated results (PDF text),
earnings presentations (FY26+), plus hand-entered values from press releases / OCR'd image-only PDFs
(see techm_handentered.py). Writes data/raw/techm.csv.

Run:  python3 scripts/extract/techm_extract.py
Requires pdftotext (poppler) and openpyxl. Source files in data/sources/techm/.
"""
import os, re, glob, subprocess, csv, sys
import openpyxl
sys.path.insert(0, os.path.dirname(__file__))
from techm_common import *
import techm_handentered as HE

os.makedirs(TXT, exist_ok=True)
LINKS = link_map()
ROWS = []
WARN = []


def ir_quarter(fn):
    """IR label 'Q1 - FY 2016-2017' -> 'Q1FY17'"""
    q, t, u = LINKS[fn]
    m = re.match(r"Q([1-4]) - FY (\d{4})-(\d{4})", q)
    return fq(int(m.group(1)), int(m.group(3))), t, u


def pdf_text(fn):
    out = os.path.join(TXT, re.sub(r"\.(pdf|PDF)$", ".txt", fn))
    if not os.path.exists(out):
        subprocess.run(["pdftotext", "-layout", os.path.join(SRC, fn), out], check=False)
    return open(out, encoding="utf-8", errors="ignore").read() if os.path.exists(out) else ""


# ---------------------------------------------------------------- doc dates (from press-release datelines)
DOC_DATE = {}
for fn in LINKS:
    q, t, u = LINKS[fn]
    if t == "Press Release" and os.path.exists(os.path.join(SRC, fn)):
        dq, _, _ = ir_quarter(fn)
        d = dateline(pdf_text(fn), after=period_end(dq))
        if d:
            DOC_DATE[dq] = d
DOC_DATE.update(HE.DOC_DATE_OVERRIDES)


def doc_date(dq):
    if dq in DOC_DATE:
        return DOC_DATE[dq]
    # approx: ~4 weeks after quarter end
    pe = period_end(dq)
    y, m, d = map(int, pe.split("-"))
    m += 1
    if m > 12:
        m, y = 1, y + 1
    return f"{y}-{m:02d}-28"


def add(*a, **k):
    ROWS.append(row(*a, **k))


# ---------------------------------------------------------------- label -> metric maps
def split_label(line):
    m = re.match(r"^\s*(.+?)\s{2,}(\S.*)$", line)
    if not m:
        return None, []
    label, rest = m.group(1).strip(), m.group(2)
    toks = rest.split()
    vals = []
    for t in toks:
        if t in ("`",):
            continue
        if not re.fullmatch(r"\(?-?[\d,]*\.?\d+\)?%?|-", t):
            return None, []
        vals.append(t)
    return label, vals


HC_DIMS = {"software professionals": "Software professionals", "bpo professionals": "BPO professionals",
           "bps professionals": "BPS professionals", "sales & support": "Sales & support",
           "total employees": "total", "total headcount": "total", "it": "Software professionals (IT)",
           "bps": "BPS professionals"}


def metric_for(section, label):
    """Return (metric, dimension, dim_type, unit, basis, period_type, note) or None."""
    L = label.lower().strip()
    if section == "hc":
        if L.startswith("it attrition"):
            return ("attrition", "total", "total", "pct", "na", "ltm",
                    "IT attrition % LTM; organic business (#); IT (software) professionals only")
        if L.startswith("it utilization % (excluding trainees)"):
            return ("utilization_excl_trainees", "total", "total", "pct", "na", "quarter",
                    "IT utilization excluding trainees; organic business (#)")
        if L.startswith("it utilization"):
            return ("utilization_incl_trainees", "total", "total", "pct", "na", "quarter",
                    "IT utilization (incl. trainees); organic business (#)")
        for k, v in HC_DIMS.items():
            if L == k:
                dt_ = "total" if v == "total" else "other"
                return ("headcount", v, dt_, "count", "na", "point",
                        "Consolidated headcount at period end incl. BPO/BPS and sales & support" if v == "total"
                        else "Headcount component as reported (period end)")
        return None
    if section in ("geo", "ind"):
        if L == "total":
            return None
        return ("revenue_share", label, "geography" if section == "geo" else "vertical", "pct", "na", "quarter", "")
    if section == "clients":
        if L.startswith("no. of active clients"):
            return ("active_clients", "total", "total", "count", "na", "point", "")
        if L.startswith("% of repeat"):
            return ("repeat_business_pct", "total", "total", "pct", "na", "quarter", "% of revenue from repeat business")
        return None
    if section == "buckets":
        m = re.match(r"≥\s*\$\s*(\d+)\s*million", label)
        if m:
            return ("clients_bucket", f"USD{m.group(1)}mn+", "client_bucket", "count", "na", "point",
                    "Number of clients >= USD X mn (TechM 'No. of Million $ Clients'; LTM revenue basis not stated in fact sheet)")
        return None
    if section == "conc":
        if L in ("top 5", "top 10", "top 20"):
            return ("client_concentration", label, "client_bucket", "pct", "na", "quarter",
                    "Revenue share of top-N clients (quarter)")
        return None
    if section == "onoff_rev":
        if L in ("onsite", "offshore"):
            return (f"revenue_share_{L}", label, "other", "pct", "na", "quarter", "Share of IT business revenue")
        return None
    if section == "onoff_hc":
        if L in ("onsite", "offshore"):
            return (f"headcount_share_{L}", label, "other", "pct", "na", "point",
                    "Share of IT headcount (replaced IT revenue on/off split from Q1FY23)")
        return None
    if section == "fx":
        if L.startswith("period average"):
            return ("fx_usdinr_avg", "total", "total", "ratio", "na", "quarter", "INR per USD, period average")
        return None
    if section == "deals":
        if L.startswith("net new deal wins") or L == "total":
            return ("tcv", "total", "total", "USD_mn", "na", "quarter", "Net new deal wins (TCV), USD mn")
        if L in ("communications", "enterprise", "communications, media & entertainment (cme)"):
            return ("tcv", label, "vertical", "USD_mn", "na", "quarter", "Net new deal wins (TCV) by business, USD mn")
        return None
    return None


# ---------------------------------------------------------------- fact sheets: old format (Q1FY17..Q3FY20)
def old_section(line):
    l = line.lower()
    if "total headcount (as at" in l: return "hc"
    if "revenue by geography" in l: return "geo"
    if "revenue by industry" in l: return "ind"
    if "no. of active clients &" in l: return "clients"
    if "no. of million $ clients" in l: return "buckets"
    if "client contribution" in l: return "conc"
    if "on/off break-up" in l: return "onoff_rev"
    if "rupee usd rate" in l: return "fx"
    if "proportion of revenues" in l or "hedge" in l: return "skip"
    if "receivable days" in l: return "skip"
    return None


def parse_old_factsheet(fn, text):
    dq, t, url = ir_quarter(fn)
    k, F = parse_fq(dq)
    prev = [fq(i, F - 1) for i in range(1, 5)]
    cur = [fq(i, F) for i in range(1, k + 1)]
    cols_tot = prev + ["T"] + cur + (["T"] if k == 4 else [])
    cols_no = prev + cur
    doc = f"TechM {dq} fact sheet"
    dd = doc_date(dq)
    mode, section, got_rev = None, None, set()
    for line in text.splitlines():
        if "P&L Summary (Rs in Mn)" in line: mode = "INR"; section = None
        if "P&L Summary (US$ in Mn)" in line: mode = "USD"; section = None
        s = old_section(line)
        if s:
            section = s
            continue
        label, vals = split_label(line)
        if not label or not vals:
            continue
        n = len(vals)
        cols = cols_tot if n == len(cols_tot) else cols_no if n == len(cols_no) else None
        if label.lower().startswith(("revenue from services", "revenue from operations")) and mode and mode not in got_rev:
            if cols is None:
                WARN.append(f"{fn}: revenue row n={n}"); continue
            got_rev.add(mode)
            for c, v in zip(cols, vals):
                if c == "T" or to_num(v) is None: continue
                add(c, "revenue", "total", "total", to_num(v), "USD_mn" if mode == "USD" else "INR_mn", "reported",
                    "quarter", url, doc, f"P&L Summary ({'US$' if mode=='USD' else 'Rs'} in Mn)", dd,
                    "Revenue from services/operations, consolidated")
            continue
        if section in (None, "skip"):
            continue
        # Top-N rows sit under 'Client Contribution' header even after the buckets block in some sheets
        sec = section
        if label.lower() in ("top 5", "top 10", "top 20"): sec = "conc"
        if label.startswith("≥"): sec = "buckets"
        mm = metric_for(sec, label)
        if not mm:
            continue
        if cols is None:
            WARN.append(f"{fn}: {label} n={n} (expected {len(cols_tot)}/{len(cols_no)})"); continue
        metric, dim, dtp, unit, basis, pt, note = mm
        for c, v in zip(cols, vals):
            if c == "T": continue
            x = to_num(v)
            if x is None: continue
            add(c, metric, dim, dtp, x, unit, basis, pt, url, doc, "Consolidated Fact Sheet Data (5+ quarter table)", dd, note)


# ---------------------------------------------------------------- fact sheets: new format (Q1FY21..Q4FY25)
QLAB = re.compile(r"Q([1-4])\s?[’']?\s?FY\s?(\d{2})")


def new_section(name):
    l = name.lower()
    if "revenue by industry" in l: return "ind"
    if "revenue by geography" in l: return "geo"
    if "on/off revenue" in l: return "onoff_rev"
    if "headcount onsite" in l: return "onoff_hc"
    if "deal wins" in l: return "deals"
    if "active clients" in l: return "clients"
    if "million $ clients" in l: return "buckets"
    if "client concentration" in l: return "conc"
    if "total headcount" in l: return "hc"
    if "attrition" in l: return "hc"
    if "usd rupee rate" in l: return "fx"
    if "p&l in inr" in l: return "pl_inr"
    if "p&l in usd" in l: return "pl_usd"
    return "skip"


def parse_new_factsheet(fn, text):
    dq, t, url = ir_quarter(fn)
    doc = f"TechM {dq} fact sheet"
    dd = doc_date(dq)
    cols, section, growth = None, None, False
    for line in text.splitlines():
        labs = QLAB.findall(line)
        if len(labs) >= 3:
            first = QLAB.search(line)
            name = line[:first.start()].strip()
            cols = [fq(int(a), 2000 + int(b)) for a, b in labs[:3]]
            section = new_section(name)
            growth = False
            continue
        if re.search(r"FY\s?\d{2}\s+FY\s?\d{2}", line) or "KEY HIGHLIGHTS" in line:
            cols, section, growth = None, None, False   # annual tables / highlight banners
            continue
        if "Revenue Growth (USD)" in line:
            growth, section = True, None
            continue
        label, vals = split_label(line)
        if not label or not vals:
            continue
        if growth:
            if len(vals) == 4 and all(v.endswith("%") for v in vals):
                is_tot = label.lower().startswith("total")
                dim, dtp = ("total", "total") if is_tot else (label, "vertical")
                met = "revenue_growth_%s" if is_tot else "segment_growth_%s"
                note = "USD revenue growth as reported in fact sheet 'Revenue Growth (USD)' table"
                for (per, basis), v in zip([("qoq", "reported"), ("qoq", "cc"), ("yoy", "reported"), ("yoy", "cc")], vals):
                    add(dq, met % per, dim, dtp, to_num(v), "pct", basis, "quarter", url, doc, "Revenue Growth (USD) table", dd, note)
            continue
        if cols is None or section in (None, "skip"):
            continue
        if section in ("pl_inr", "pl_usd"):
            if label.lower().startswith("revenue from operations") and len(vals) >= 3:
                for c, v in zip(cols, vals[:3]):
                    add(c, "revenue", "total", "total", to_num(v), "INR_mn" if section == "pl_inr" else "USD_mn",
                        "reported", "quarter", url, doc, "P&L in %s Mn" % ("INR" if section == "pl_inr" else "USD"), dd,
                        "Revenue from operations, consolidated")
            continue
        mm = metric_for(section, label)
        if not mm or len(vals) < 3:
            continue
        metric, dim, dtp, unit, basis, pt, note = mm
        for c, v in zip(cols, vals[:3]):
            x = to_num(v)
            if x is None: continue
            add(c, metric, dim, dtp, x, unit, basis, pt, url, doc, f"{section} table (3-quarter)", dd, note)
        if section in ("ind", "geo") and len(vals) >= 5:
            for per, v in zip(("qoq", "yoy"), vals[3:5]):
                add(cols[0], f"segment_growth_{per}", dim, dtp, to_num(v), "pct", "reported", "quarter", url, doc,
                    f"Revenue by {'Industry' if section=='ind' else 'Geography'} % table, {per.upper()} column", dd,
                    "Segment revenue growth as shown next to mix table (USD reported terms presumed; fact sheet does not label currency)")


# ---------------------------------------------------------------- consolidated results (SEBI format)
def first_nums(rest, n=3):
    toks = rest.split()
    out = []
    for t in toks:
        if re.fullmatch(r"\(?-?[\d,]*\.?\d+\)?", t):
            out.append(to_num(t))
        elif t == "-":
            out.append(None)
        else:
            if out: break
        if len(out) == n: break
    return out


def parse_results(fn, text):
    dq, t, url = ir_quarter(fn)
    doc = f"TechM {dq} consolidated financial results (SEBI format)"
    dd = doc_date(dq)
    head = text[:4000]
    if re.search(r"Lakhs", head, re.I):
        scale, unote = 0.1, "reported in Rs lakh; converted to INR mn (/10)"
    elif re.search(r"Million", head, re.I):
        scale, unote = 1.0, ""
    else:
        WARN.append(f"{fn}: unit not found"); return
    cols = [dq, prev_q(dq), prev_q(dq, 4)]
    found = set()
    seg_mode = False
    for line in text.splitlines():
        l = line.strip()
        ll = l.lower()
        if "standalone information" in ll:
            break_after = True
        if "segment revenue" in ll and "seg" not in found:
            seg_mode = True; continue
        spec = None
        if re.match(r"^(\d+\s+)?(revenue from operations|income from operations \(net\))", ll) and "rev" not in found:
            spec = ("rev", "revenue", "total", "total", "Revenue from operations (consolidated)")
        elif re.match(r"^(\w\)\s*)?(subcontracting expense|services rendered by business associates)", ll) and "sub" not in found:
            spec = ("sub", "subcontracting_cost", "total", "total",
                    "Consolidated 'Subcontracting Expense(s)' line" + ("" if "subcontract" in ll else
                    " (labelled 'Services rendered by Business Associates and Others' in this pre-Ind-AS-format filing)"))
        elif re.match(r"^(\w\)\s*)?employee benefits? expense", ll) and "emp" not in found:
            spec = ("emp", "employee_cost", "total", "total", "Employee benefits expense (consolidated)")
        elif seg_mode and re.match(r"^a\)\s*it\b", ll):
            spec = ("segit", "segment_revenue", "IT", "service_line", "Primary business segment revenue (IT services)")
        elif seg_mode and re.match(r"^b\)\s*(bpo|bps)\b", ll):
            lab = "BPS" if "bps" in ll else "BPO"
            spec = ("segbp", "segment_revenue", lab, "service_line", f"Primary business segment revenue ({lab})")
        if not spec:
            continue
        key, metric, dim, dtp, note = spec
        m = re.match(r"^(?:\d+\s+)?(?:\w\)\s*)?[A-Za-z][A-Za-z /&()\-,.]*?\s{2,}(.*)$", l)
        if not m:
            continue
        nums = first_nums(m.group(1), 3)
        if len(nums) < 3:
            WARN.append(f"{fn}: {metric} only {nums}"); continue
        found.add(key)
        if key == "segbp":
            seg_mode = False; found.add("seg")
        for c, v in zip(cols, nums):
            if v is None: continue
            val = round(v * scale, 1)
            add(c, metric, dim, dtp, val, "INR_mn", "reported", "quarter", url, doc,
                "Statement of consolidated results / segment information, quarter-ended columns", dd,
                "; ".join(x for x in [note, unote] if x))
    for k in ("rev", "sub", "emp", "segit", "segbp"):
        if k not in found:
            WARN.append(f"{fn}: missing {k}")


# ---------------------------------------------------------------- datasheets (xlsx), FY17+ columns
DS_BLOCKS = [
    ("revenue by industry % (quarter ended) -restated", "ind_restated"),
    ("revenue by industry % (quarter ended)", "ind"),
    ("on/off break-up", "onoff_rev"),
    ("it headcount onsite", "onoff_hc"),
    ("revenue by geography", "geo"),
    ("total headcount  (as at", "hc"), ("total headcount (as at", "hc"),
    ("no. of active clients &", "clients"),
    ("no. of million $ clients", "buckets"),
    ("client concentration", "conc"),
    ("deal wins (usd mn) -restated", "deals_r"),
    ("deal wins (usd mn)", "deals"),
    ("cash flows", "skip"), ("rupee usd rate", "fx"), ("proportion of revenues", "skip"),
    ("hedge book", "skip"), ("notes", "skip"),
]
IND_CLASS = {1: "industry classification v1 (to FY21; 'Discontinued in FY22')",
             2: "industry classification v2 (CME/Technology; FY22-FY23, restated FY21)",
             3: "industry classification v3 (Communications/Hi-Tech & Media/HLS; FY24+, restated FY23)"}


def parse_datasheet(fn, full_history):
    dq, t, url = ir_quarter(fn)
    doc = f"TechM {dq} data sheet (xlsx)"
    dd = doc_date(dq)
    wb = openpyxl.load_workbook(os.path.join(SRC, fn), data_only=True)
    keep = None if full_history else set([dq] + [prev_q(dq, i) for i in range(1, 5)])

    def colmap(ws):
        cm, fy = {}, None
        for j in range(2, ws.max_column + 1):
            a = ws.cell(1, j).value
            if isinstance(a, str) and re.search(r"FY\s*(\d{4})-(\d{2})", a):
                mm = re.search(r"FY\s*(\d{4})-(\d{2})", a)
                fy = 2000 + int(mm.group(2))
            b = ws.cell(2, j).value
            if fy and isinstance(b, str) and re.fullmatch(r"Q[1-4]", b.strip()) and fy >= 2017:
                q = fq(int(b.strip()[1]), fy)
                if keep is None or q in keep:
                    cm[j] = q
        return cm

    names = {n.strip().lower(): n for n in wb.sheetnames}
    # revenue
    for key, unit in (("p&l rs mn", "INR_mn"), ("p&l us$ mn", "USD_mn")):
        if key not in names: WARN.append(f"{fn}: no sheet {key}"); continue
        ws = wb[names[key]]
        cm = colmap(ws)
        for r in range(3, 12):
            lab = ws.cell(r, 1).value
            if isinstance(lab, str) and lab.strip().lower().startswith(("revenue from operations", "revenue from services")):
                for j, q in cm.items():
                    v = ws.cell(r, j).value
                    if isinstance(v, (int, float)):
                        add(q, "revenue", "total", "total", round(v, 2), unit, "reported", "quarter", url, doc,
                            f"sheet '{names[key]}', row '{lab.strip()}'", dd, "Revenue from operations, consolidated (unrounded xlsx value)")
                break
    ws = wb[names["operating metrics"]]
    cm = colmap(ws)
    block, ind_ver, deal_ver = None, 0, 0
    for r in range(3, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if not isinstance(lab, str):
            continue
        L = lab.strip().lower()
        hit = None
        for k, b in DS_BLOCKS:
            if L.startswith(k):
                hit = b; break
        if hit:
            block = hit
            if hit in ("ind", "ind_restated"):
                ind_ver += 1
            if hit in ("deals", "deals_r"):
                deal_ver += 1
            continue
        if L.startswith("*") or L.startswith("as part of"):
            continue
        sec = {"ind_restated": "ind", "deals_r": "deals"}.get(block, block)
        if sec == "hc" and L.startswith(("it attrition", "it utilization")):
            pass
        mm = metric_for(sec, lab.strip()) if sec not in (None, "skip") else None
        if not mm:
            continue
        metric, dim, dtp, unit, basis, pt, note = mm
        if sec == "ind":
            note = (note + "; " if note else "") + IND_CLASS.get(ind_ver, "")
        if sec == "deals" and deal_ver == 1:
            note += "; original (pre-FY22) Communications/Enterprise split"
        if sec == "deals" and deal_ver == 2:
            note += "; restated series (CME/Enterprise split to FY23; total only thereafter)"
        for j, q in cm.items():
            v = ws.cell(r, j).value
            if not isinstance(v, (int, float)):
                continue
            if unit == "pct":
                v = round(v * 100, 3)
            elif unit == "count":
                v = round(v)
            else:
                v = round(v, 3)
            blk = f" [industry classification v{ind_ver}]" if sec == "ind" else (f" [deal-wins block {deal_ver}]" if sec == "deals" else "")
            add(q, metric, dim, dtp, v, unit, basis, pt, url, doc, f"sheet 'Operating Metrics', row '{lab.strip()}'{blk}", dd,
                note + "; xlsx unrounded value (pct stored as fraction, x100)" if unit == "pct" else note)


# ---------------------------------------------------------------- earnings presentations FY26+ (geo/vertical mix & growth)
def parse_presentation(fn, text):
    dq, t, url = ir_quarter(fn)
    doc = f"TechM {dq} earnings presentation"
    dd = doc_date(dq)
    i = text.find("Geography-wise")
    if i < 0:
        i = text.find("Geography-wise".lower())
    seg = text[i:i + 4000] if i >= 0 else ""
    mode = None
    for line in seg.splitlines():
        s = line.strip()
        if s.startswith("Geographies"): mode = "geography"; continue
        if s.startswith("Verticals"): mode = "vertical"; continue
        if s.startswith("Copyright") or s.startswith("Client Metrics"): break
        m = re.match(r"^([A-Za-z][A-Za-z ,&]+?)\s+(-?[\d.]+)%\s+(-?[\d.]+)%\s+(-?[\d.]+)%", s)
        if m and mode:
            dim = m.group(1).strip()
            note = "As reported in earnings presentation 'Geography-wise and Vertical-wise Performance' (growth presumably USD reported)"
            add(dq, "revenue_share", dim, mode, float(m.group(2)), "pct", "na", "quarter", url, doc, "Geography-wise and Vertical-wise Performance", dd, note)
            add(dq, "segment_growth_qoq", dim, mode, float(m.group(3)), "pct", "reported", "quarter", url, doc, "Geography-wise and Vertical-wise Performance", dd, note)
            add(dq, "segment_growth_yoy", dim, mode, float(m.group(4)), "pct", "reported", "quarter", url, doc, "Geography-wise and Vertical-wise Performance", dd, note)


# ---------------------------------------------------------------- run
def main():
    for fn in sorted(LINKS):
        path = os.path.join(SRC, fn)
        if not os.path.exists(path):
            continue
        q, t, u = LINKS[fn]
        tl = t.lower().replace(" ", "")
        if tl in ("factsheet",) and fn.lower().endswith(".pdf"):
            text = pdf_text(fn)
            if len(text.strip()) < 100:
                WARN.append(f"{fn}: image-only PDF (fact sheet) - not parsed"); continue
            dq, _, _ = ir_quarter(fn)
            k, F = parse_fq(dq)
            if F <= 2020:
                parse_old_factsheet(fn, text)
            else:
                parse_new_factsheet(fn, text)
        elif tl in ("consolidatedresult", "consolidatedresults"):
            text = pdf_text(fn)
            if len(text.strip()) < 100:
                WARN.append(f"{fn}: image-only PDF (results) - hand-entered from OCR where needed"); continue
            parse_results(fn, text)
        elif tl == "datasheet":
            parse_datasheet(fn, full_history=(fn == "tml-q1-fy-27-datasheet.xlsx"))
        elif tl in ("earningpresentation", "earningspresentation"):
            dq, _, _ = ir_quarter(fn)
            if parse_fq(dq)[1] >= 2026:
                parse_presentation(fn, pdf_text(fn))
    # hand-entered (press releases, OCR of image PDFs)
    for h in HE.rows(LINKS, doc_date):
        ROWS.append(h)
    # tag industry classification on fact-sheet / presentation vertical rows (by document fiscal year)
    for r in ROWS:
        if r["dim_type"] == "vertical" and r["metric"] in ("revenue_share", "segment_growth_qoq", "segment_growth_yoy") \
                and ("fact sheet" in r["source_doc"] or "presentation" in r["source_doc"]):
            dq = r["source_doc"].split()[1]
            q_, fy_ = parse_fq(dq)
            v = 1 if fy_ <= 2021 else 2 if fy_ <= 2023 else 3
            tag = IND_CLASS[v]
            if dq == "Q4FY25":
                tag += "; Q4FY25 sheet: customers with multiple businesses re-aligned to verticals, prior-year comparative re-aligned"
            r["notes"] = (r["notes"] + "; " if r["notes"] else "") + tag
    out = os.path.join(ROOT, "data/raw/techm.csv")
    # de-duplicate exact duplicates (same doc/quarter/metric/dim/basis/unit/period_type)
    seen, final = set(), []
    for r in ROWS:
        key = (r["fiscal_q"], r["metric"], r["dimension"], r["basis"], r["unit"], r["period_type"], r["source_url"], r["source_loc"])
        if key in seen:
            continue
        seen.add(key)
        final.append(r)
    final.sort(key=lambda r: (r["period_end"], r["metric"], r["dimension"], r["doc_date"], r["source_url"]))
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(final)
    print(f"wrote {len(final)} rows to {out}")
    print("\n".join(WARN))


if __name__ == "__main__":
    main()
